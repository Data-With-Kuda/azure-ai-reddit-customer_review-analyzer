import os
import praw
import pandas as pd
from datetime import datetime
from dotenv import load_dotenv
from azure.ai.textanalytics import TextAnalyticsClient
from azure.core.credentials import AzureKeyCredential
from openai import AzureOpenAI
from docx import Document

load_dotenv()

# Environment Variables
REDDIT_CLIENT_ID       = os.getenv("REDDIT_CLIENT_ID")
REDDIT_CLIENT_SECRET   = os.getenv("REDDIT_CLIENT_SECRET")
REDDIT_USERNAME        = os.getenv("REDDIT_USERNAME")
REDDIT_PASSWORD        = os.getenv("REDDIT_PASSWORD")
REDDIT_USER_AGENT      = os.getenv("REDDIT_USER_AGENT", "windows:reddit_analyzer:0.1 (by u/KNyam)")

AZURE_LANGUAGE_KEY     = os.getenv("AZURE_LANGUAGE_KEY")
AZURE_LANGUAGE_ENDPOINT= os.getenv("AZURE_LANGUAGE_ENDPOINT")

AZURE_OPENAI_KEY       = os.getenv("AZURE_OPENAI_KEY")
AZURE_OPENAI_ENDPOINT  = os.getenv("AZURE_OPENAI_ENDPOINT")
AZURE_OPENAI_DEPLOYMENT = os.getenv("AZURE_OPENAI_DEPLOYMENT")

# 1. Fetch Reddit Comments
def fetch_reddit_comments(post_id: str) -> pd.DataFrame:
    reddit = praw.Reddit(
        client_id=REDDIT_CLIENT_ID,
        client_secret=REDDIT_CLIENT_SECRET,
        username=REDDIT_USERNAME,
        password=REDDIT_PASSWORD,
        user_agent=REDDIT_USER_AGENT
    )
    submission = reddit.submission(id=post_id)
    submission.comments.replace_more(limit=0)
    data = [{
        "author": str(c.author),
        "text": c.body,
        "score": c.score,
        "created_utc": datetime.utcfromtimestamp(c.created_utc).isoformat()
    } for c in submission.comments.list()]
    return pd.DataFrame(data)

# 2. Azure NLP Analysis
def azure_nlp_analysis(df: pd.DataFrame) -> pd.DataFrame:
    client = TextAnalyticsClient(
        endpoint=AZURE_LANGUAGE_ENDPOINT,
        credential=AzureKeyCredential(AZURE_LANGUAGE_KEY)
    )
    results = []
    batch_size = 10
    texts = df["text"].tolist()
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i+batch_size]
        sentiments = client.analyze_sentiment(documents=batch)
        key_phrases = client.extract_key_phrases(documents=batch)
        for s, k in zip(sentiments, key_phrases):
            phrases = "" if k.is_error else ", ".join(k.key_phrases)
            results.append({
                "sentiment": s.sentiment,
                "positive_score": s.confidence_scores.positive,
                "neutral_score": s.confidence_scores.neutral,
                "negative_score": s.confidence_scores.negative,
                "key_phrases": phrases
            })
    return pd.DataFrame(results)

# 3. Generate Individual Replies
def generate_responses(texts: list[str]) -> list[str]:
    client = AzureOpenAI(
        api_key=AZURE_OPENAI_KEY,
        api_version="2023-05-15",
        azure_endpoint=AZURE_OPENAI_ENDPOINT
    )
    responses = []
    for text in texts:
        try:
            resp = client.chat.completions.create(
                model=AZURE_OPENAI_DEPLOYMENT,
                messages=[
                    {"role": "system", "content": "You are a friendly customer support assistant."},
                    {"role": "user", "content": f"Write a polite reply to: '{text}'"}
                ],
                temperature=0.6
            )
            responses.append(resp.choices[0].message.content)
        except Exception as e:
            responses.append(f"Error: {e}")
    return responses

# 4. Summarize Key Themes & Recommendations
def summarize_key_phrases(key_phrases: list[str]) -> str:
    client = AzureOpenAI(
        api_key=AZURE_OPENAI_KEY,
        api_version="2023-05-15",
        azure_endpoint=AZURE_OPENAI_ENDPOINT
    )
    lines = "\n".join(f"- {kp}" for kp in key_phrases if kp)
    prompt = (
        "You are a product insights specialist. Given the following key phrases from Reddit feedback, "
        "please provide:\n1. A concise summary of the top 3 recurring themes.\n"
        "2. For each theme, suggest 1–2 actionable product improvements.\n\n"
        f"{lines}"
    )
    resp = client.chat.completions.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        messages=[
            {"role": "system", "content": "You summarize feedback into themes and actionable recommendations."},
            {"role": "user", "content": prompt}
        ],
        temperature=0.5,
        max_tokens=500
    )
    return resp.choices[0].message.content

# 5. Generate Word Report
def generate_word_report(content: str, filename: str):
    doc = Document()
    doc.add_heading("🎯 Reddit Feedback: Themes & Product Recommendations", level=1)
    doc.add_paragraph(f"Report created: {datetime.now():%Y-%m-%d %H:%M}\n")
    doc.add_paragraph(content)
    doc.save(filename)

# Main Pipeline
def run_pipeline(reddit_post_url: str):
    post_id = reddit_post_url.split("/comments/")[1].split("/")[0]
    print("▶️  Fetching comments...")
    df = fetch_reddit_comments(post_id)

    print("🔍 Analyzing sentiment and key phrases...")
    nlp_df = azure_nlp_analysis(df)

    print("📝 Generating AI replies...")
    df["ai_response"] = generate_responses(df["text"].tolist())

    result = pd.concat([df, nlp_df], axis=1)

    # Extract all key phrases
    all_phrases = []
    for kp in result["key_phrases"].dropna():
        all_phrases.extend([p.strip() for p in kp.split(",") if p.strip()])

    print("📊 Summarizing key themes and recommendations...")
    summary = summarize_key_phrases(all_phrases)

    # Output files
    base = f"reddit_analysis_{post_id}_{datetime.now():%Y%m%d_%H%M%S}"
    result.to_csv(base + ".csv", index=False)
    generate_word_report(summary, base + ".docx")

    print("✅ Done! Outputs:")
    print(" - Detailed CSV:", base + ".csv")
    print(" - Management Report:", base + ".docx")

if __name__ == "__main__":
    run_pipeline("https://www.reddit.com/r/cocacola/comments/1j5g5ib/opinions_on_this_one/")
